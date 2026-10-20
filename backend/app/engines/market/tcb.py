"""TCB daily retail price lists for Dhaka (WP-09 milestone 1): XLSX rows -> PricePoints.

Each list (one XLSX per publication day, saved under `data/public/prices/tcb/`) gives, per item,
TCB's lowest and highest retail price today, one week before and one month before. Every one of
those three columns becomes a PricePoint dated by its own header date, with the midpoint of the
range as the price. When two lists give the same date (a list's own "today" and a later list's
"one week before"), the list published on that date wins, then the newest list.

Nothing here guesses: an item missing from a list, or given as 0-0 (not on sale), adds no point.
Item names are matched after Unicode NFD normalisation, because TCB's files mix the precomposed
and decomposed forms of letters such as য় and ড়.
"""

from __future__ import annotations

import csv
import re
import unicodedata
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from app.contracts import DataMode, Money, PricePoint, PriceSource

REPO_ROOT = Path(__file__).resolve().parents[4]
TCB_DIR = REPO_ROOT / "data" / "public" / "prices" / "tcb"
LISTING_URL = "https://tcb.gov.bd/pages/daily-rmps"
MARKET = "Dhaka"

CENT = Decimal("0.01")
_PRICE_COLUMNS = ((2, 3), (4, 5), (6, 7))  # (low, high) for today, 1 week before, 1 month before
_COLUMN_NAMES = ("today", "week_before", "month_before")


@dataclass(frozen=True)
class Item:
    commodity: str
    name_bn: str  # prefix of TCB's item name
    unit: str


ITEMS: tuple[Item, ...] = (
    Item("rice_coarse", "চাল (মোটা)", "kg"),
    Item("rice_fine", "চাল সরু", "kg"),
    Item("atta_loose", "আটা সাদা (খোলা)", "kg"),
    Item("soybean_oil", "সয়াবিন তেল (লুজ)", "litre"),
    Item("lentil_masur", "মশুর ডাল (বড় দানা)", "kg"),
    Item("chickpea", "ছোলা", "kg"),
    Item("sugar", "চিনি", "kg"),
    Item("onion", "পিঁয়াজ (দেশী)", "kg"),
    Item("potato", "আলু", "kg"),
    Item("egg_farm", "ডিম (ফার্ম)", "hali"),
    Item("garlic", "রসুন (দেশী)", "kg"),
    Item("salt", "লবণ", "kg"),
    Item("ms_rod", "এম,এস রড (৬০", "metric_ton"),
)


@dataclass(frozen=True)
class Observation:
    """A price point and where it was read: the list it came from and the row/column."""

    point: PricePoint
    ref: str  # e.g. "tcb:2026-10-19.xlsx#r12:today"
    published_on: date


class TcbFormatError(ValueError):
    """The list does not have TCB's layout (no header with the three dated price columns)."""


def _norm(value: object) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFD", str(value or ""))).strip()


_ITEMS = tuple((_norm(item.name_bn), item) for item in ITEMS)
_HEADER = _norm("অদ্যকার")  # "today's price"
_END = _norm("যে সকল")  # the notes below the price table


def _money(value: Decimal) -> Money:
    return Money.of(value.quantize(CENT, rounding=ROUND_HALF_UP))


def _number(value: object) -> Decimal | None:
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        return None
    return Decimal(str(value))


def _header_dates(rows: Sequence[Sequence[object]]) -> tuple[int, tuple[date, date, date]]:
    for index, row in enumerate(rows[:-1]):
        if len(row) > 2 and _HEADER in _norm(row[2]):
            dated = rows[index + 1]
            cells = [dated[low] if len(dated) > low else None for low, _ in _PRICE_COLUMNS]
            if all(isinstance(cell, datetime | date) for cell in cells):
                days = tuple(c.date() if isinstance(c, datetime) else c for c in cells)
                return index + 2, days  # type: ignore[return-value]
    raise TcbFormatError("no header row with the dates of the three price columns")


def parse_list(rows: Sequence[Sequence[object]], file_name: str, published_on: date,
               source_url: str | None = None) -> list[Observation]:
    """The watched items of one TCB list (its rows as read from the sheet)."""
    first, days = _header_dates(rows)
    found: list[Observation] = []
    seen: set[str] = set()
    for number, row in enumerate(rows[first:], start=first + 1):
        name = _norm(row[0] if row else None)
        if name.startswith(_END):
            break
        item = next((it for prefix, it in _ITEMS if name.startswith(prefix)), None)
        if item is None or item.commodity in seen:
            continue
        seen.add(item.commodity)
        for (low_col, high_col), day, column in zip(_PRICE_COLUMNS, days, _COLUMN_NAMES,
                                                     strict=True):
            low = _number(row[low_col]) if len(row) > high_col else None
            high = _number(row[high_col]) if len(row) > high_col else None
            if low is None or high is None or low <= 0 or high < low:
                continue  # not on sale that day, or not a range: no point
            point = PricePoint(
                commodity=item.commodity, unit=item.unit, price=_money((low + high) / 2),
                low=_money(low), high=_money(high), market=MARKET, source=PriceSource.tcb,
                source_url=source_url, observed_on=day, data_mode=DataMode.public)
            found.append(Observation(point, f"tcb:{file_name}#r{number}:{column}", published_on))
    return found


def merge(observations: Iterable[Observation]) -> list[Observation]:
    """One observation per (commodity, date): the list published that day, else the newest."""
    best: dict[tuple[str, date], Observation] = {}
    for obs in observations:
        key = (obs.point.commodity, obs.point.observed_on)
        current = best.get(key)
        if current is None or _rank(obs) > _rank(current):
            best[key] = obs
    return sorted(best.values(), key=lambda o: (o.point.commodity, o.point.observed_on))


def _rank(obs: Observation) -> tuple[bool, date]:
    return obs.published_on == obs.point.observed_on, obs.published_on


def load_directory(directory: Path = TCB_DIR) -> list[Observation]:
    """Every saved list in `directory` (`<YYYY-MM-DD>.xlsx`, URLs in `sources.csv`), merged."""
    from openpyxl import load_workbook  # heavy import, only when lists are read

    urls: dict[str, str] = {}
    sources = directory / "sources.csv"
    if sources.exists():
        with sources.open(encoding="utf-8", newline="") as handle:
            urls = {row["date"]: row["url"] for row in csv.DictReader(handle)}
    observations: list[Observation] = []
    for path in sorted(directory.glob("*.xlsx")):
        published = date.fromisoformat(path.stem)
        sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
        observations += parse_list(rows, path.name, published, urls.get(path.stem, LISTING_URL))
    return merge(observations)
