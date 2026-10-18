"""Read an uploaded CSV or XLSX file into raw rows (WP-02).

Real exports are messy: a byte-order mark, UTF-16 "Unicode text" saved by Excel, semicolon or
tab delimiters, account details above the table and totals below it. This module only turns
bytes into rows of cell text with a stable reference to where each row came from; finding the
header and reading the cells is the parsers' job.
"""

from __future__ import annotations

import codecs
import csv
import io
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal

MAX_ROWS = 200_000
_DELIMITERS = ",;\t|"


class TableError(ValueError):
    """The file cannot be read as a table at all."""

    def __init__(self, reason_key: str) -> None:
        super().__init__(reason_key)
        self.reason_key = reason_key


@dataclass(frozen=True)
class Row:
    ref: str
    cells: tuple[str, ...]

    @property
    def blank(self) -> bool:
        return not any(cell.strip() for cell in self.cells)


@dataclass(frozen=True)
class Sheet:
    name: str | None
    rows: tuple[Row, ...]


def is_xlsx(data: bytes) -> bool:
    return data[:4] == b"PK\x03\x04"


def decode_text(data: bytes) -> str:
    """Decode CSV bytes: BOM first, then UTF-8, then Latin-1 (never fails, never guesses Bangla)."""
    if data.startswith(codecs.BOM_UTF8):
        return data[len(codecs.BOM_UTF8):].decode("utf-8")
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=_DELIMITERS).delimiter
    except csv.Error:
        counts = {d: sample.count(d) for d in _DELIMITERS}
        best = max(counts, key=lambda d: counts[d])
        return best if counts[best] else ","


def read_csv(data: bytes, filename: str) -> Sheet:
    text = decode_text(data)
    lines = text.splitlines()
    sample = "\n".join(lines[:50])
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=_delimiter(sample))
    rows: list[Row] = []
    try:
        for cells in reader:
            if len(rows) >= MAX_ROWS:
                raise TableError("ingest.file.too_many_rows")
            rows.append(Row(ref=f"{filename}#L{reader.line_num}",
                            cells=tuple(cell.strip() for cell in cells)))
    except csv.Error as exc:
        raise TableError("ingest.file.malformed_csv") from exc
    return Sheet(name=None, rows=tuple(rows))


def cell_text(value: object) -> str:
    """Text for one spreadsheet cell. Floats keep their shortest exact decimal form."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, float):
        number = Decimal(repr(value))
        return format(number.normalize(), "f") if number == number.to_integral() else repr(value)
    return str(value).strip()


def read_xlsx(data: bytes, filename: str) -> list[Sheet]:
    from openpyxl import load_workbook  # imported lazily: only spreadsheet uploads need it

    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise TableError("ingest.file.malformed_xlsx") from exc
    sheets = []
    try:
        for ws in book.worksheets:
            rows: list[Row] = []
            for number, values in enumerate(ws.iter_rows(values_only=True), start=1):
                if len(rows) >= MAX_ROWS:
                    raise TableError("ingest.file.too_many_rows")
                rows.append(Row(ref=f"{filename}!{ws.title}#R{number}",
                                cells=tuple(cell_text(v) for v in values)))
            sheets.append(Sheet(name=ws.title, rows=tuple(rows)))
    finally:
        book.close()
    return sheets


def read_sheets(data: bytes, filename: str) -> list[Sheet]:
    """Every sheet of the file (a CSV is one sheet)."""
    if not data:
        raise TableError("ingest.file.empty")
    if is_xlsx(data):
        return read_xlsx(data, filename)
    return [read_csv(data, filename)]
